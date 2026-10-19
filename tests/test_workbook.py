from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import load_workbook

from workbook import collect, create_template, save_atomic


class WorkbookTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        self.source = self.root / "input.xlsx"
        self.output = self.root / "output/results.xlsx"
        create_template(self.source)

    def rows(self, *rows):
        book = load_workbook(self.source)
        for row in rows:
            book.active.append(row)
        book.save(self.source)
        book.close()

    def value(self, path, cell):
        book = load_workbook(path)
        try:
            return book.active[cell].value
        finally:
            book.close()

    def lookup(self, phone, directory):
        directory.mkdir(parents=True, exist_ok=True)
        avatar = directory / f"{phone}.png"
        avatar.write_bytes(b"test avatar")
        return "测试昵称", avatar

    def test_resume_and_retry_missing_avatar_without_changing_input(self):
        self.rows(("13800000000",))
        before = self.source.read_bytes()
        lookup = Mock(side_effect=self.lookup)
        collect(self.source, self.output, lookup, interval=0)
        self.assertEqual(self.value(self.output, "B2"), "测试昵称")
        avatar = self.output.parent / self.value(self.output, "C2")
        self.assertTrue(avatar.is_file())
        collect(self.source, self.output, lookup, interval=0)
        self.assertEqual(lookup.call_count, 1)
        avatar.unlink()
        collect(self.source, self.output, lookup, interval=0)
        self.assertEqual(lookup.call_count, 2)
        self.assertEqual(self.source.read_bytes(), before)

    def test_failure_stops_and_preserves_completed_rows(self):
        self.rows(("13800000000",), ("13900000000",), ("13700000000",))
        calls = []

        def lookup(phone, directory):
            calls.append(phone)
            if len(calls) == 2:
                raise RuntimeError("window closed")
            return self.lookup(phone, directory)

        with self.assertRaisesRegex(RuntimeError, "window closed"):
            collect(self.source, self.output, lookup, interval=0)
        self.assertEqual(len(calls), 2)
        self.assertEqual(self.value(self.output, "B2"), "测试昵称")
        self.assertIsNone(self.value(self.output, "B3"))

    def test_empty_results_retry_and_invalid_phones_skip(self):
        self.rows(("13800000000",), ("not-a-phone",), (None,))
        lookup = Mock(return_value=(None, None))
        collect(self.source, self.output, lookup, interval=0)
        collect(self.source, self.output, lookup, interval=0)
        self.assertEqual(lookup.call_count, 2)
        self.assertIsNone(self.value(self.output, "B2"))

    def test_rejects_same_path_and_bad_headers(self):
        with self.assertRaises(ValueError):
            collect(self.source, self.source, Mock())
        book = load_workbook(self.source)
        book.active["A1"] = "unexpected"
        book.save(self.source)
        book.close()
        lookup = Mock()
        with self.assertRaises(ValueError):
            collect(self.source, self.output, lookup)
        lookup.assert_not_called()
        self.assertFalse(self.output.exists())

    def test_init_does_not_overwrite(self):
        before = self.source.read_bytes()
        with self.assertRaises(FileExistsError):
            create_template(self.source)
        self.assertEqual(self.source.read_bytes(), before)

    def test_failed_replace_preserves_previous_workbook(self):
        before = self.source.read_bytes()
        book = load_workbook(self.source)
        book.active["A2"] = "new data"
        try:
            with patch.object(Path, "replace", side_effect=PermissionError("file locked")):
                with self.assertRaises(PermissionError):
                    save_atomic(book, self.source)
        finally:
            book.close()
        self.assertEqual(self.source.read_bytes(), before)
        self.assertEqual(list(self.root.glob("*.xlsx")), [self.source])

    def test_nickname_is_text_not_formula(self):
        self.rows(("13800000000",))
        collect(self.source, self.output, Mock(return_value=("=nickname", None)), interval=0)
        book = load_workbook(self.output)
        try:
            self.assertEqual(book.active["B2"].value, "=nickname")
            self.assertEqual(book.active["B2"].data_type, "s")
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
