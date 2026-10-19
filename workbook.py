"""表格读写和逐行采集，不依赖 Windows 桌面。"""

import logging
from pathlib import Path
import re
import tempfile
import time

from openpyxl import Workbook, load_workbook


COLUMNS = ("phone_number", "nickname", "avatar_file")


def save_atomic(book, path):
    """先完整写入临时文件，再替换结果，避免中途退出损坏旧结果。"""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".xlsx", delete=False) as temp:
        temporary = Path(temp.name)
    try:
        book.save(temporary)
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def create_template(path):
    path = Path(path)
    if path.exists():
        raise FileExistsError(f"文件已存在：{path}")
    book = Workbook()
    book.active.append(COLUMNS)
    save_atomic(book, path)
    book.close()


def text(value):
    return str(value).strip() if value is not None else ""


def collect(source, output, lookup, interval=3):
    source, output = Path(source).resolve(), Path(output).resolve()
    if source.suffix.lower() != ".xlsx" or output.suffix.lower() != ".xlsx":
        raise ValueError("输入和输出文件必须为 .xlsx")
    if source == output:
        raise ValueError("输出文件必须与输入文件不同")
    if not source.is_file():
        raise FileNotFoundError(f"找不到输入表格：{source}，可先用 --init 创建")
    book = load_workbook(output if output.exists() else source)
    try:
        sheet = book.active
        if tuple(sheet.cell(1, col).value for col in range(1, 4)) != COLUMNS:
            raise ValueError("前三列表头必须为 phone_number、nickname、avatar_file")
        avatars = output.parent / "avatars"
        save_atomic(book, output)
        for row in range(2, sheet.max_row + 1):
            phone, nickname, avatar = (text(sheet.cell(row, col).value) for col in range(1, 4))
            if not phone:
                continue
            if not re.fullmatch(r"1[3-9]\d{9}", phone, flags=re.ASCII):
                logging.warning("第 %s 行：手机号格式无效，跳过", row)
                continue
            if nickname and avatar and (output.parent / avatar).is_file():
                continue
            logging.info("查询第 %s 行（尾号 %s）", row, phone[-4:])
            found_name, found_avatar = lookup(phone, avatars)
            if found_name:
                cell = sheet.cell(row, 2, found_name)
                cell.data_type = "s"  # 昵称以 = 开头时仍保存为文本。
            if found_avatar:
                sheet.cell(row, 3, Path(found_avatar).resolve().relative_to(output.parent).as_posix())
            save_atomic(book, output)
            if not found_name or not found_avatar:
                logging.warning("第 %s 行：结果不完整，下次运行会重试", row)
            time.sleep(interval)
        logging.info("处理结束：%s", output)
    finally:
        book.close()
